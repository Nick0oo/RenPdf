"""Live workbook acceptance: invented XLSX/PDF only, with --cloud opt-in."""

import argparse
import io
import json
import re
import time
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import openpyxl
import pymupdf
from dotenv import dotenv_values

ROOT = Path(__file__).resolve().parents[1]


def workbook_bytes(first: bool) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Presupuesto" if first else "Revisión"
    sheet.append(["Proyecto", "Presupuesto USD"])
    sheet.append(["Aurora", 120 if first else 70])
    sheet["B2"].number_format = "$#,##0.00"
    if first:
        sheet = workbook.create_sheet("Ajuste")
        sheet.append(["Proyecto", "Presupuesto USD"])
        sheet.append(["Aurora", 90])
        sheet = workbook.create_sheet("Fórmulas")
        sheet.append(["Concepto", "Resultado"])
        sheet.append(["Suma guardada", "=SUM(Presupuesto!B2,Ajuste!B2)"])
        sheet.append(["Sin resultado guardado", "=B2*10"])
    output = io.BytesIO()
    workbook.save(output)
    if not first:
        return output.getvalue()
    patched = io.BytesIO()
    with ZipFile(output) as original, ZipFile(patched, "w", ZIP_DEFLATED) as target:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet3.xml":
                data = re.sub(
                    rb'(<c r="B2"[^>]*>.*?<v>).*?(</v>)', rb"\g<1>210\g<2>", data
                )
            target.writestr(entry, data)
    return patched.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cloud", action="store_true")
    parser.add_argument("--url", default="http://localhost:8000")
    parser.add_argument("--resume", action="store_true")
    args = parser.parse_args()
    if not args.cloud:
        parser.error("--cloud authorizes sending only the invented fixtures")
    credentials = dotenv_values(ROOT / ".env")
    report_path = ROOT / "docs" / "spreadsheet-rag-results.json"
    result = (
        json.loads(report_path.read_text(encoding="utf-8"))
        if args.resume and report_path.is_file()
        else {
            "date": datetime.now(UTC).isoformat(),
            "synthetic_data_only": True,
            "checks": {},
        }
    )
    client = httpx.Client(base_url=args.url + "/api/v1", timeout=180)

    def save():
        report_path.write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    def request(method, path, **kwargs):
        response = client.request(method, path, **kwargs)
        response.raise_for_status()
        return response

    def settled(identifier):
        retries, deadline = 0, time.monotonic() + 480
        while time.monotonic() < deadline:
            jobs = request("GET", f"/documents/{identifier}/jobs").json()
            if any(job["state"] in {"QUEUED", "RUNNING"} for job in jobs):
                time.sleep(1)
                continue
            if jobs and jobs[0]["state"] != "SUCCEEDED":
                if (
                    jobs[0]["error_code"]
                    in {"PROVIDER_UNAVAILABLE", "PROVIDER_TIMEOUT"}
                    and retries < 2
                ):
                    retries += 1
                    request(
                        "POST", f"/documents/{identifier}/jobs/{jobs[0]['id']}/retry"
                    )
                    continue
                raise RuntimeError(
                    f"Synthetic job failed: {jobs[0]['kind']} / {jobs[0]['error_code']}"
                )
            document = request("GET", f"/documents/{identifier}").json()
            assert document["current_version_id"]
            return document
        raise RuntimeError("Workbook acceptance processing timeout")

    def grant(identifier, scope):
        request(
            "POST",
            f"/documents/{identifier}/consents",
            json={"provider": "gemini", "scope": scope, "expires_in_days": 1},
        )

    def ask(key, identifiers, question):
        for attempt in range(3):
            response = request(
                "POST",
                f"/conversations/{result['conversation_id']}/messages",
                json={"question": question, "document_ids": identifiers},
            )
            if "event: done" in response.text:
                break
            if attempt == 2 or not any(
                code in response.text
                for code in ("PROVIDER_UNAVAILABLE", "PROVIDER_TIMEOUT")
            ):
                raise RuntimeError(response.text)
        answer = request(
            "GET", f"/conversations/{result['conversation_id']}/messages"
        ).json()[-1]
        assert answer["state"] == "COMPLETE" and answer["citations"], answer
        context = request("GET", f"/trace/responses/{answer['id']}/context").json()
        assert context["query"] == question and context["sources"]
        for citation in answer["citations"]:
            if citation["source_type"] == "spreadsheet":
                assert (
                    citation["page_number"] is None
                    and citation["sheet_name"]
                    and citation["cell_range"]
                )
        result["checks"][key] = {
            "id": answer["id"],
            "question": question,
            "answer": answer["content"],
            "citations": answer["citations"],
            "retrieval_mode": answer["retrieval_mode"],
            "trace_context_count": len(context["sources"]),
        }
        save()
        print(f"Answered {key}: {answer['content']}", flush=True)
        return answer

    try:
        request(
            "POST",
            "/login/access-token",
            data={
                "username": credentials["FIRST_SUPERUSER"],
                "password": credentials["FIRST_SUPERUSER_PASSWORD"],
            },
        )
        catalog = request("GET", "/providers/").json()
        models = {
            model["id"]
            for provider in catalog
            if provider["name"] == "gemini"
            for model in provider["models"]
        }
        routing = request("GET", "/providers/routing").json()
        model = next(
            (
                row["model"]
                for row in routing
                if row["task"] == "chat"
                and row["provider"] == "gemini"
                and row["model"] in models
            ),
            None,
        )
        model = model or next(
            value
            for value in [credentials.get("DEFAULT_CHAT_MODEL"), "gemini-2.5-flash"]
            if value in models
        )
        result["chat_model"] = model
        if not args.resume:
            workspace = request(
                "POST",
                "/workspaces/",
                json={
                    "name": "Excel · hojas, celdas y comparación",
                    "description": "Libros y acta inventados para verificar valores, fórmulas y citas exactas.",
                },
            ).json()
            result["workspace_id"] = workspace["id"]
            result["sources"] = []
            save()
            with pymupdf.open() as pdf:
                pdf.new_page().insert_text(
                    (40, 50),
                    "Acta Aurora: el presupuesto aprobado en este documento es de 80 dolares.",
                )
                payloads = [
                    ("Aurora-presupuesto.xlsx", workbook_bytes(True)),
                    ("Aurora-revision.xlsx", workbook_bytes(False)),
                    ("Aurora-acta.pdf", pdf.tobytes()),
                ]
            for title, payload in payloads:
                document = request(
                    "POST",
                    f"/documents/?workspace_id={result['workspace_id']}",
                    files={"file": (title, payload)},
                ).json()
                identifier = document["id"]
                request(
                    "PUT",
                    f"/workspaces/{result['workspace_id']}/documents/{identifier}",
                )
                settled(identifier)
                grant(identifier, "chat")
                coverage = request("GET", f"/documents/{identifier}/coverage").json()
                if document["source_type"] == "spreadsheet":
                    assert coverage["sheets_total"] == coverage["sheets_extracted"]
                    assert coverage["cells_total"] == coverage["cells_with_chunks"]
                trace = request("GET", f"/trace/documents/{identifier}").json()
                active = [
                    index["profile"]
                    for version in trace["versions"]
                    for index in version["indexes"]
                    if index["active"]
                ]
                assert active and active[0]["provider"] in {"llamacpp", "ollama"}
                result["sources"].append(
                    {
                        "id": identifier,
                        "title": title,
                        "source_type": document["source_type"],
                        "local_profile": active[0],
                        "local_coverage": coverage,
                    }
                )
                save()
                print(f"Prepared and locally indexed {title}", flush=True)
            conversation = request(
                "POST",
                f"/workspaces/{result['workspace_id']}/conversations",
                json={"provider": "gemini", "model": model},
            ).json()
            result["conversation_id"] = conversation["id"]
            save()
        identifiers = [source["id"] for source in result["sources"]]
        primary = identifiers[0]
        answer = ask(
            "exact_cell",
            [primary],
            "¿Qué valor guardado tiene la celda B2 de la hoja Ajuste? Cita la celda.",
        )
        assert "90" in answer["content"]
        assert all(
            citation["sheet_name"] == "Ajuste" and citation["cell_range"] == "B2"
            for citation in answer["citations"]
        )
        answer = ask(
            "cross_sheet",
            [primary],
            "Compara el presupuesto de Aurora en las hojas Presupuesto y Ajuste: valores guardados y diferencia. Cita ambas hojas.",
        )
        assert all(value in answer["content"] for value in ("120", "90", "30"))
        assert {citation["sheet_name"] for citation in answer["citations"]} == {
            "Presupuesto",
            "Ajuste",
        }
        answer = ask(
            "cross_documents",
            identifiers,
            "Compara los presupuestos de Aurora: hojas Presupuesto y Ajuste del primer libro, hoja Revisión del segundo libro y el acta PDF. Indica las cuatro cifras y cita cada origen.",
        )
        assert all(value in answer["content"] for value in ("120", "90", "70", "80"))
        assert {citation["document_id"] for citation in answer["citations"]} == set(
            identifiers
        )
        answer = ask(
            "cached_formula",
            [primary],
            "¿Qué resultado GUARDADO tiene la fórmula de la celda B2 en la hoja Fórmulas? Cita la celda.",
        )
        assert "210" in answer["content"] and all(
            citation["cell_range"] == "B2" for citation in answer["citations"]
        )
        answer = ask(
            "uncached_formula",
            [primary],
            "¿Existe un resultado GUARDADO en la celda B3 de la hoja Fórmulas? No calcules la fórmula, verifica si existe valor guardado y cita la celda.",
        )
        assert any(
            term in answer["content"].lower()
            for term in ("no ", "sin ", "no disponible")
        )
        assert all(citation["cell_range"] == "B3" for citation in answer["citations"])
        model = next(
            value
            for value in [
                credentials.get("DEFAULT_EMBEDDING_MODEL"),
                "gemini-embedding-2",
            ]
            if value in models and "embedding" in value
        )
        result["gemini_embedding_model"] = model
        for identifier in identifiers[:2]:
            grant(identifier, "embeddings")
            request(
                "POST",
                f"/documents/{identifier}/index",
                json={
                    "provider": "gemini",
                    "model": model,
                    "modality": "text",
                    "dimensions": 768,
                },
            )
            settled(identifier)
        answer = ask(
            "gemini_text_embeddings",
            identifiers[:2],
            "Compara los valores de Aurora en las hojas Presupuesto, Ajuste y Revisión de los dos libros. Indica las tres cifras y cita cada hoja.",
        )
        assert (
            all(value in answer["content"] for value in ("120", "90", "70"))
            and answer["retrieval_mode"] == "hybrid"
        )
        assert {citation["document_id"] for citation in answer["citations"]} == set(
            identifiers[:2]
        )
        sources = request(
            "GET",
            f"/trace/rag/sources?workspace_id={result['workspace_id']}&source_type=spreadsheet",
        ).json()
        queries = request(
            "GET", f"/trace/rag/queries?workspace_id={result['workspace_id']}"
        ).json()
        assert sources["total"] == 2 and queries["total"] >= 6
        chunks = request(
            "GET", f"/trace/documents/{primary}/chunks?sheet=Fórmulas&cell=B3"
        ).json()["data"]
        assert chunks and all(item["analysis"]["cached"] is False for item in chunks)
        result["checks"]["trace_module"] = {
            "workbooks": sources["total"],
            "queries": queries["total"],
            "uncached_formula_inspectable": True,
        }
        result["passed"] = True
        result.pop("error", None)
        result["finished_at"] = datetime.now(UTC).isoformat()
        print("All live workbook acceptance checks passed", flush=True)
    except Exception as error:
        result["passed"], result["error"] = False, str(error)
        raise
    finally:
        save()
        client.close()


if __name__ == "__main__":
    main()
