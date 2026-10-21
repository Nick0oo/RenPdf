"""Real Excel extraction, exact cell citations and mixed-format RAG."""

import io
import json
import uuid
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pymupdf
import pytest
import xlwt
from sqlmodel import Session, col, select

from app.core.db import engine
from app.domain import Document, DocumentElement, DocumentPage, Job
from app.services.models.contracts import (
    ModelCapabilities,
    ModelDescriptor,
    ProviderError,
)
from app.services.parsing.spreadsheet import read_spreadsheet
from app.services.rag.retrieval import retrieve
from app.workers.tasks import process_job

FIXTURE = Path(__file__).parent / "fixtures" / "spreadsheet_workspace.xlsx"


def upload_and_extract(client, headers, payload, filename):
    response = client.post(
        "/api/v1/documents/", headers=headers, files={"file": (filename, payload)}
    )
    assert response.status_code == 200, response.text
    identifier = response.json()["id"]
    with Session(engine) as session:
        job = session.exec(
            select(Job).where(col(Job.document_id) == uuid.UUID(identifier))
        ).one()
        process_job(str(job.id))
        session.expire_all()
        assert session.get(Job, job.id).state == "SUCCEEDED"
    return identifier


@pytest.fixture
def workbook_source(client, superuser_token_headers, monkeypatch):
    monkeypatch.setattr("app.api.routes.documents.enqueue", lambda _job: None)
    identifier = upload_and_extract(
        client,
        superuser_token_headers,
        FIXTURE.read_bytes(),
        f"workbook-{uuid.uuid4()}.xlsx",
    )
    yield identifier
    client.delete(f"/api/v1/documents/{identifier}", headers=superuser_token_headers)


def test_real_workbook_preserves_values_formulas_merges_and_worksheet_locations():
    book = read_spreadsheet(FIXTURE)
    assert [sheet.name for sheet in book.sheets] == ["Ventas", "Resumen"]
    cells = {cell.coordinate: cell for cell in book.sheets[0].cells}
    assert cells["D2"].value == 30 and cells["D2"].cached
    assert cells["D2"].formula == "=SUM(B2:C2)"
    assert cells["E2"].value is None and not cells["E2"].cached
    assert cells["G3"].merged_range == "G3:I3"
    assert cells["B40"].value == 400
    assert book.warnings and "sin resultado" in book.warnings[0]


def test_local_extraction_and_paginated_cell_view_are_owner_scoped(
    workbook_source, client, superuser_token_headers, normal_user_token_headers
):
    identifier = workbook_source
    document = client.get(
        f"/api/v1/documents/{identifier}", headers=superuser_token_headers
    ).json()
    assert document["source_type"] == "spreadsheet" and document["sheet_count"] == 2
    assert document["page_count"] == 0 and document["current_version_id"]
    assert (
        client.get(
            f"/api/v1/documents/{identifier}/consents", headers=superuser_token_headers
        ).json()
        == []
    )
    sheets = client.get(
        f"/api/v1/documents/{identifier}/sheets", headers=superuser_token_headers
    )
    assert sheets.status_code == 200, sheets.text
    assert [sheet["name"] for sheet in sheets.json()["data"]] == ["Ventas", "Resumen"]
    cells = client.get(
        f"/api/v1/documents/{identifier}/sheets/cells?sheet=Ventas&row=39&limit=3&column=2&columns=2",
        headers=superuser_token_headers,
    ).json()["cells"]
    assert any(cell["coordinate"] == "B40" and cell["value"] == 400 for cell in cells)
    assert all(39 <= cell["row"] <= 41 and 2 <= cell["column"] <= 3 for cell in cells)
    coverage = client.get(
        f"/api/v1/documents/{identifier}/coverage", headers=superuser_token_headers
    ).json()
    assert coverage["sheets_total"] == coverage["sheets_extracted"] == 2
    assert coverage["cells_total"] == coverage["cells_with_chunks"] == 269
    assert coverage["missing_formula_values"] == 1
    for suffix in ("sheets", "sheets/cells?sheet=Ventas", "coverage"):
        assert (
            client.get(
                f"/api/v1/documents/{identifier}/{suffix}",
                headers=normal_user_token_headers,
            ).status_code
            == 404
        )
    assert (
        client.get(
            f"/api/v1/documents/{identifier}/sheets/cells?sheet=Inexistente",
            headers=superuser_token_headers,
        ).status_code
        == 404
    )
    assert (
        client.get(
            f"/api/v1/documents/{identifier}/pages/1/elements",
            headers=superuser_token_headers,
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/v1/documents/{identifier}/analyze",
            headers=superuser_token_headers,
            json={"provider": "gemini", "model": "test"},
        ).status_code
        == 422
    )
    assert (
        client.post(
            f"/api/v1/documents/{identifier}/index",
            headers=superuser_token_headers,
            json={
                "provider": "gemini",
                "model": "gemini-embedding-2",
                "modality": "multimodal",
            },
        ).status_code
        == 422
    )


def test_exact_cell_rag_citations_context_export_and_trace_module(
    workbook_source,
    client,
    superuser_token_headers,
    normal_user_token_headers,
    monkeypatch,
):
    identifier = workbook_source
    with Session(engine) as session:
        both, _ = retrieve(
            session,
            [uuid.UUID(identifier)],
            "Compara las celdas B2 y C2 de la hoja Ventas",
        )
        assert {source.cell_range for source in both} == {"B2", "C2"}
        sources, _mode = retrieve(
            session,
            [uuid.UUID(identifier)],
            "¿Qué valor tiene la celda B3 de la hoja Resumen?",
        )
        assert sources and all(
            source.sheet_name == "Resumen" and source.cell_range == "B3"
            for source in sources
        )
        assert all(
            source.page_number is None and source.precision == "cell"
            for source in sources
        )

    class Provider:
        is_cloud = False

        def discover(self):
            return [
                ModelDescriptor(
                    id="cell-test",
                    provider="ollama",
                    version="test",
                    capabilities=ModelCapabilities(text="declared"),
                )
            ]

        def stream(self, request):
            evidence = json.loads(request.prompt)["evidence"]
            assert evidence and all(
                item["sheet_name"] == "Resumen" and item["cell_range"] == "B3"
                for item in evidence
            )
            assert all(item["page_number"] is None for item in evidence)
            yield f"El valor guardado es 30. [Fuente {evidence[0]['source_id']}]"

    monkeypatch.setattr("app.services.rag.chat.get_provider", lambda *_args: Provider())
    workspace = client.post(
        "/api/v1/workspaces/",
        headers=superuser_token_headers,
        json={"name": "Exact spreadsheet cells"},
    ).json()
    try:
        client.put(
            f"/api/v1/workspaces/{workspace['id']}/documents/{identifier}",
            headers=superuser_token_headers,
        )
        conversation = client.post(
            f"/api/v1/workspaces/{workspace['id']}/conversations",
            headers=superuser_token_headers,
            json={"provider": "ollama", "model": "cell-test"},
        ).json()
        question = "¿Qué valor tiene la celda B3 de la hoja Resumen?"
        response = client.post(
            f"/api/v1/conversations/{conversation['id']}/messages",
            headers=superuser_token_headers,
            json={"question": question, "document_ids": [identifier]},
        )
        assert "event: done" in response.text, response.text
        answer = client.get(
            f"/api/v1/conversations/{conversation['id']}/messages",
            headers=superuser_token_headers,
        ).json()[-1]
        assert (
            answer["citations"][0]["sheet_name"] == "Resumen"
            and answer["citations"][0]["cell_range"] == "B3"
        )
        context = client.get(
            f"/api/v1/trace/responses/{answer['id']}/context",
            headers=superuser_token_headers,
        ).json()
        assert context["query"] == question and all(
            item["cell_range"] == "B3" for item in context["sources"]
        )
        sources = client.get(
            f"/api/v1/trace/rag/sources?source_type=spreadsheet&workspace_id={workspace['id']}",
            headers=superuser_token_headers,
        )
        assert sources.status_code == 200 and sources.json()["total"] == 1
        chunks = client.get(
            f"/api/v1/trace/documents/{identifier}/chunks?sheet=Ventas&cell=E2",
            headers=superuser_token_headers,
        ).json()["data"]
        assert len(chunks) == 1 and chunks[0]["analysis"]["cached"] is False
        assert (
            chunks[0]["analysis"]["value"] is None and chunks[0]["cell_range"] == "E2"
        )
        assert (
            client.get(
                f"/api/v1/trace/documents/{identifier}/chunks?cell=BAD",
                headers=superuser_token_headers,
            ).status_code
            == 422
        )
        exported = client.get(
            f"/api/v1/conversations/{conversation['id']}/export",
            headers=superuser_token_headers,
        ).text
        assert "hoja Resumen, celda B3" in exported
        assert (
            client.get(
                f"/api/v1/trace/documents/{identifier}/chunks?sheet=Ventas",
                headers=normal_user_token_headers,
            ).status_code
            == 404
        )
    finally:
        client.delete(
            f"/api/v1/workspaces/{workspace['id']}", headers=superuser_token_headers
        )


def test_cross_sheet_and_cross_document_retrieval_keeps_the_short_sheet(
    client, superuser_token_headers, monkeypatch
):
    monkeypatch.setattr("app.api.routes.documents.enqueue", lambda _job: None)
    identifiers = []
    try:
        for first in (True, False):
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Ventas" if first else "Corrección"
            sheet.append(["Proyecto", "Presupuesto"])
            sheet.append(["Aurora", 120 if first else 70])
            if first:
                for index in range(50):
                    sheet.append([f"Otro proyecto {index}", index])
                summary = workbook.create_sheet("Resumen")
                summary.append(["Proyecto", "Presupuesto"])
                summary.append(["Aurora", 90])
            output = io.BytesIO()
            workbook.save(output)
            identifiers.append(
                upload_and_extract(
                    client,
                    superuser_token_headers,
                    output.getvalue(),
                    f"budget-{first}.xlsx",
                )
            )
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((40, 50), "Aurora presupuesto PDF 80 pesos")
            identifiers.append(
                upload_and_extract(
                    client, superuser_token_headers, pdf.tobytes(), "budget.pdf"
                )
            )
        with Session(engine) as session:
            sources, _mode = retrieve(
                session,
                [uuid.UUID(value) for value in identifiers],
                "Compara el presupuesto de Aurora entre ambos libros y el PDF",
            )
            assert {str(source.document_id) for source in sources} == set(identifiers)
            assert {
                source.sheet_name
                for source in sources
                if source.source_type == "spreadsheet"
            } == {"Ventas", "Resumen", "Corrección"}
            assert all(
                source.page_number is None
                for source in sources
                if source.source_type == "spreadsheet"
            )
            assert any(
                source.cell_range == "B2" and "90" in source.text for source in sources
            )
            scoped, _mode = retrieve(
                session,
                [uuid.UUID(identifiers[0])],
                "Compara el presupuesto de Aurora entre las hojas Ventas y Resumen",
            )
            assert {source.sheet_name for source in scoped} == {"Ventas", "Resumen"}
    finally:
        for identifier in identifiers:
            client.delete(
                f"/api/v1/documents/{identifier}", headers=superuser_token_headers
            )


def test_cancellation_retries_from_the_next_sheet_without_publishing_partial_cells(
    client, superuser_token_headers, monkeypatch
):
    monkeypatch.setattr("app.api.routes.documents.enqueue", lambda _job: None)
    response = client.post(
        "/api/v1/documents/",
        headers=superuser_token_headers,
        files={"file": ("checkpoint.xlsx", FIXTURE.read_bytes())},
    )
    identifier = response.json()["id"]
    from app.workers.tasks import check_cancel

    def stop_after_first(session, job):
        check_cancel(session, job)
        if job.checkpoint == 1:
            raise ProviderError("CANCELLED", "Synthetic sheet checkpoint")

    try:
        monkeypatch.setattr("app.workers.tasks.check_cancel", stop_after_first)
        with Session(engine) as session:
            job = session.exec(
                select(Job).where(col(Job.document_id) == uuid.UUID(identifier))
            ).one()
            process_job(str(job.id))
            session.expire_all()
            assert session.get(Job, job.id).checkpoint == 1
            assert session.get(Job, job.id).state == "CANCELLED"
            assert (
                session.get(Document, uuid.UUID(identifier)).current_version_id is None
            )
        monkeypatch.setattr("app.workers.tasks.check_cancel", check_cancel)
        assert (
            client.post(
                f"/api/v1/documents/{identifier}/jobs/{job.id}/retry",
                headers=superuser_token_headers,
            ).status_code
            == 200
        )
        process_job(str(job.id))
        with Session(engine) as session:
            document = session.get(Document, uuid.UUID(identifier))
            units = session.exec(
                select(DocumentPage).where(
                    col(DocumentPage.version_id) == document.current_version_id
                )
            ).all()
            assert len(units) == 2 and all(unit.unit_kind == "sheet" for unit in units)
            assert (
                len(
                    session.exec(
                        select(DocumentElement)
                        .join(DocumentPage)
                        .where(
                            col(DocumentPage.version_id) == document.current_version_id
                        )
                    ).all()
                )
                == 269
            )
    finally:
        client.delete(
            f"/api/v1/documents/{identifier}", headers=superuser_token_headers
        )


def test_legacy_xls_and_utf8_csv_are_read_without_executing_formulas(tmp_path):
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Legado")
    sheet.write(0, 0, "Proyecto")
    sheet.write(1, 0, "Aurora")
    sheet.write(1, 1, 123)
    path = tmp_path / "legacy.xls"
    workbook.save(str(path))
    book = read_spreadsheet(path)
    assert book.format == "xls" and book.sheets[0].cells[-1].value == 123
    path = tmp_path / "data.csv"
    path.write_text(
        'Código,Importe,Nota\n001,12.5,"=HYPERLINK(https://invalid)"\n',
        encoding="utf-8",
    )
    book = read_spreadsheet(path, path.name)
    cells = {cell.coordinate: cell for cell in book.sheets[0].cells}
    assert cells["A2"].value == "001" and cells["B2"].value == 12.5
    assert cells["C2"].formula is None and cells["C2"].kind == "text"


def test_limits_use_actual_cells_not_untrusted_dimensions_and_reject_non_workbooks(
    tmp_path, monkeypatch
):
    path = tmp_path / "bounded.xlsx"
    with ZipFile(FIXTURE) as original, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:R65"', b'ref="A1:XFD1048576"')
            target.writestr(entry, data)
    assert read_spreadsheet(path).sheets[0].rows == 65
    monkeypatch.setattr(
        "app.services.parsing.spreadsheet.settings.MAX_SPREADSHEET_CELLS", 5
    )
    with pytest.raises(ValueError, match="SPREADSHEET_CELL_LIMIT"):
        read_spreadsheet(path)
    with ZipFile(path, "w") as target:
        target.writestr("word/document.xml", "Not an Excel workbook")
    with pytest.raises(ValueError, match="INVALID_SPREADSHEET"):
        read_spreadsheet(path)
